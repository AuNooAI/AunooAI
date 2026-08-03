"""Contract tests: every renderer must consume the shapes the generators
actually produce.

Three of the 2026-08-03 Q3 review defects were silent key mismatches — a
renderer reading ``name``/``trajectory``/``timeframe`` where the generator
writes ``title``/``description``/``time_horizon``, so every card printed an
em dash; and the exec-letter payload reading ``impact_score`` where the EOS
generator writes ``impact_rating``, so the letter claimed "no new tail-risk
scenarios" while the deck showed 24 cards. These tests push each
generator's REAL output shape through each renderer and assert the content
survives. No DB, no LLM — pure shapes.
"""
import re

import pytest


# ── Fixtures: the real generator shapes ───────────────────────────────

def eos_card() -> dict:
    """Exactly OutlierScenario.to_dict — the EOS generator's shape."""
    from app.services.extreme_outlier_service import OutlierScenario
    return OutlierScenario(
        id="eos-1",
        category="wild_card",
        title="Collapse of Scientific Consensus",
        subtitle="sub",
        description="A cascade of retractions [44, 52] erodes trust [11].",
        probability=0.08,
        impact_rating=9,
        time_horizon="2028-2032",
    ).to_dict()


def exec_summary_card() -> dict:
    """The prompt-v3 exec summary card shape (no invented percentages)."""
    return {
        "topic_title": "CITATION FRAUD & REFERENCE INTEGRITY",
        "primary_horizon": "h1",
        "horizon_label": "Declining System",
        "opening_statement": "Most scenarios expect fabricated references to "
                             "become a routine editorial problem [3].",
        "minority_view": {"statement": "Verification tooling could outpace abuse."},
        "primary_signal": "Reference verification becomes a submission "
                          "requirement at major publishers [7][12].",
        "decision_fork": {
            "condition_a": {"condition": "If verification ships", "outcome": "A"},
            "condition_b": {"condition": "If it slips", "outcome": "B"},
        },
        "action_window": {
            "assessment": {"timeframe": "0-6 months", "action": "assess"},
            "positioning": {"timeframe": "6-18 months", "action": "position"},
        },
        "source_scenarios": [{"title": "Scenario One", "horizon": "h1"}],
    }


def articles(n: int = 60) -> list:
    return [{"title": f"Article {i}", "uri": f"https://example.com/{i}",
             "source": "example.com", "date": "2026-07-01"} for i in range(1, n + 1)]


# ── Black swan cards ──────────────────────────────────────────────────

def test_black_swan_html_renders_generator_shape():
    from app.services.topic_report_html import _render_black_swans
    html = _render_black_swans([eos_card()], 1, articles=articles())
    assert "Collapse of Scientific Consensus" in html
    assert "<h4>—</h4>" not in html
    assert "2028-2032" in html                      # time_horizon consumed
    assert "cascade of retractions" in html         # description consumed
    # grouped citation split and linked
    assert "[44]" in html and "[52]" in html and "[44, 52]" not in html


def test_black_swan_docx_renders_generator_shape():
    from docx import Document
    from app.services.topic_report_docx import _render_black_swans
    doc = Document()
    _render_black_swans(doc, [eos_card()], articles())
    text = "\n".join(p.text for p in doc.paragraphs)
    assert "Collapse of Scientific Consensus" in text
    assert "2028-2032" in text
    assert not re.search(r"^—$", text, re.M)
    assert "[44][52]" in text


def test_exec_letter_payload_counts_all_eos_categories():
    from app.services.wiley_bundle_supervisor import _exec_summary_payload
    cards = [eos_card(),
             {**eos_card(), "category": "black_swan", "impact_rating": 7,
              "title": "Second"},
             {**eos_card(), "category": "contrarian", "impact_rating": 6,
              "title": "Third"}]
    p = _exec_summary_payload([], "Q3 2026", {}, {"Topic A": cards})
    assert p["black_swan_count"] == 3               # every category counts
    assert p["top_black_swan"]["impact"] == 9       # impact_rating consumed
    assert p["top_black_swan"]["timeframe"] == "2028-2032"


# ── Exec summary cards ────────────────────────────────────────────────

def test_exec_card_html_renders_without_percentages():
    from app.services.topic_report_html import _render_executive_summary
    html = _render_executive_summary([exec_summary_card()], 1, articles=articles())
    assert "CITATION FRAUD" in html
    assert "Most scenarios expect" in html
    assert "CONSENSUS" not in html
    assert "undefined" not in html


def test_exec_card_shared_html_renders_without_percentages():
    from app.services.html_report_common import render_executive_summary_cards
    html = render_executive_summary_cards([exec_summary_card()], articles())
    assert "CITATION FRAUD" in html
    assert "CONSENSUS" not in html


def test_exec_card_docx_renders_without_percentages():
    from docx import Document
    from app.services.topic_report_docx import _render_exec_summary_cards
    doc = Document()
    _render_exec_summary_cards(doc, [exec_summary_card()], articles())
    text = "\n".join(p.text for p in doc.paragraphs)
    assert "CITATION FRAUD" in text
    assert "CONSENSUS" not in text


def test_exec_card_legacy_percentage_is_not_rendered():
    """Cached pre-v3 cards carry the invented fields — they must be ignored,
    and prose-embedded shares must be scrubbed."""
    from app.services.topic_report_html import _render_executive_summary
    legacy = exec_summary_card()
    legacy["consensus_percentage"] = 82
    legacy["minority_view"]["percentage_range"] = "7-12%"
    legacy["opening_statement"] = ("With 80% of sources expecting decline, "
                                   "pressure mounts.")
    html = _render_executive_summary([legacy], 1, articles=articles())
    assert "82" not in html
    assert "7-12%" not in html
    assert "80% of sources" not in html
    assert "most sources" in html.lower()


# ── Three Horizons scenarios ──────────────────────────────────────────

def test_scenario_html_renders_prompt_shape():
    from app.services.topic_report_html import _render_scenarios
    scenario = {"type": "h1", "title": "Fraud Becomes Routine",
                "description": "Editors adapt [2, 9].",
                "timeframe": "2026-2030", "sentiment": "Negative"}
    html = _render_scenarios([scenario], "Topic", 1, articles=articles())
    assert "Fraud Becomes Routine" in html
    assert "[2]" in html and "[9]" in html and "[2, 9]" not in html


# ── Release lint end-to-end sanity ───────────────────────────────────

def test_lint_catches_seeded_defects():
    from app.services.report_lint import lint_artifact_text
    bad = ("model: gpt-5.4\nPERSONA\n82% CONSENSUS\n"
           "with 80% of sources expecting\n[44, 52]\nby Q3 2025\n—")
    checks = {f["check"] for f in lint_artifact_text(bad, kind="deck")}
    assert checks >= {"config_leak", "invented_consensus", "grouped_citation",
                      "past_deadline", "fallback_title"}


def test_lint_passes_clean_text():
    from app.services.report_lint import lint_artifact_text
    ok = ("Most scenarios expect reference verification to ship in 2027. "
          "Retraction volume rose [4][7]. The business model: AI-assisted "
          "review remains contested.")
    assert lint_artifact_text(ok, kind="deck") == []
