"""Deterministic release lint for customer-facing report artifacts.

Runs at the end of every topic-report build. Each check encodes a defect
the 2026-08-03 Q3 review actually found; the whole point is that the next
one is caught by a machine before a human ships it. Checks are regex and
key lookups only — no LLM calls, so the lint adds milliseconds, never
minutes.

Findings are advisory: they are logged, written into the period's state
sidecar, and surfaced to the caller. Blocking the send path is the review
flow's decision, not this module's.
"""
from __future__ import annotations

import logging as _logging
import re as _re
from datetime import date as _date

_log = _logging.getLogger(__name__)


# Model/config identifiers that must never reach a customer artifact as
# OUR text. Source headlines may legitimately contain "GPT-5.6", so these
# only fire on configuration-shaped contexts ("model: gpt-5.4", a bare
# alias on a stat plate), not on any substring hit.
_CONFIG_LEAK_RES = [
    # Only fire when the value looks like an actual model id — prose like
    # "business model: AI-driven" must not trip this.
    _re.compile(r"\bmodel\s*[:=]\s*(?:gpt|claude|bedrock|nova|gemini|mixtral|kimi|gemma)[a-z0-9.\-:]*",
                _re.IGNORECASE),
    _re.compile(r"\bPERSONA\b"),
    _re.compile(r"\bCORPUS SCANNED\b"),
    _re.compile(r"\b(?:bedrock-[a-z0-9.\-]+|nova-(?:lite|pro)|litellm)\b"),
    # Internal payload/field names narrated into customer prose — the Q3
    # letter wrote "With no events_by_topic records supplied" because the
    # agent's instructions name the field.
    _re.compile(r"\b(?:events_by_topic|raw_output|per_topic_detail|"
                r"briefing_lede|eos_per_topic|scenario event file)\b"),
]

# The invented-consensus shapes prompt v3 abolished.
_CONSENSUS_RES = [
    _re.compile(r"\d{1,3}\s*%\s*CONSENSUS\b", _re.IGNORECASE),
    _re.compile(r"\d{1,3}(?:\s*[-–]\s*\d{1,3})?\s*%\s+of\s+(?:sources|scenarios)\b",
                _re.IGNORECASE),
]

_GROUPED_CITE_RE = _re.compile(r"\[\d{1,3},\s*\d{1,3}")


def _past_deadline_re() -> _re.Pattern:
    """"by Q3 2025" / "by 2024" — any named deadline before the current year."""
    this_year = _date.today().year
    past = "|".join(str(y) for y in range(2020, this_year))
    return _re.compile(rf"\bby\s+(?:Q[1-4]\s+)?(?:{past})\b", _re.IGNORECASE)


def lint_artifact_text(text: str, *, kind: str,
                       scope_topics: list | None = None,
                       workspace_topics: list | None = None) -> list:
    """String checks over a rendered artifact's full text.

    ``kind`` is "deck" or "html" (a couple of checks differ). Returns a
    list of ``{check, detail}`` findings.
    """
    findings: list = []
    if not text:
        return findings

    for rx in _CONFIG_LEAK_RES:
        for m in rx.finditer(text):
            findings.append({"check": "config_leak",
                             "detail": m.group(0)[:80]})
    for rx in _CONSENSUS_RES:
        for m in rx.finditer(text):
            findings.append({"check": "invented_consensus",
                             "detail": m.group(0)[:80]})
    for m in _GROUPED_CITE_RE.finditer(text):
        findings.append({"check": "grouped_citation", "detail": m.group(0)})
    for m in _past_deadline_re().finditer(text):
        findings.append({"check": "past_deadline", "detail": m.group(0)})

    # Em-dash fallback titles — a renderer read a key the generator never
    # writes and printed the placeholder.
    if kind == "html":
        n = text.count("<h4>—</h4>")
        if n:
            findings.append({"check": "fallback_title",
                             "detail": f"{n} card title(s) rendered as an em dash"})
    else:
        n = len(_re.findall(r"^—$", text, _re.M))
        if n:
            findings.append({"check": "fallback_title",
                             "detail": f"{n} text frame(s) containing only an em dash"})

    # Topics from OTHER customers' monitoring leaking onto this artifact.
    if workspace_topics and scope_topics:
        scope_l = {t.lower() for t in scope_topics}
        for wt in workspace_topics:
            if not wt or wt.lower() in scope_l:
                continue
            if wt.lower() in text.lower():
                findings.append({"check": "workspace_leak", "detail": wt})

    return findings


def lint_run_content(raw_output: dict, corpus: list, *, topic: str = "") -> list:
    """Content checks over one forecast run: every figure and organisation
    the prose asserts must exist in the numbered corpus it cites.

    ``corpus`` rows carry title (+ optionally summary). Findings only —
    the build is never rewritten here.
    """
    import json as _json
    findings: list = []
    if not isinstance(raw_output, dict):
        return findings

    # Body prose = the narrative fields, NOT the reference titles.
    parts: list = []
    brief = raw_output.get("topic_briefing") or {}
    parts += [brief.get("lede"), brief.get("intelligence_view"), brief.get("headline")]
    parts += [(t or {}).get("body") for t in (brief.get("tensions") or [])
              if isinstance(t, dict)]
    for s in raw_output.get("scenarios") or []:
        if isinstance(s, dict):
            parts.append(s.get("description"))
    for r in raw_output.get("strategic_recommendations") or []:
        if isinstance(r, dict):
            parts.append(r.get("rationale"))
    parts += [k for k in (raw_output.get("key_insights") or []) if isinstance(k, str)]
    body = "\n".join(p for p in parts if p)

    source_text = " ".join(
        f"{(a.get('title') or '')} {(a.get('summary') or '')}"
        for a in (corpus or []) if isinstance(a, dict)
    )
    if not body or not source_text:
        return findings

    try:
        from app.services.wiley_humanizer import (
            _FIGURE_RE, _figure_key, _figure_ledger, _org_candidates,
        )
    except Exception as e:  # never break a build over the lint's imports
        _log.warning("report lint: humanizer helpers unavailable: %s", e)
        return findings

    ledger = _figure_ledger([source_text])
    for m in _FIGURE_RE.finditer(body):
        if _figure_key(m.group(0)) not in ledger:
            findings.append({"check": "unsourced_figure",
                             "detail": f"{topic}: {m.group(0)}"})

    hay = source_text.lower()
    for key, name in _org_candidates(body).items():
        if key in hay:
            continue
        words = [w for w in _re.split(r"[^A-Za-z]+", name) if len(w) > 3]
        if any(w.lower() in hay for w in words):
            continue
        findings.append({"check": "unsourced_org",
                         "detail": f"{topic}: {name}"})
    return findings


def lint_topic_report(items: list, *, deck_text: str = "",
                      html_text: str = "",
                      workspace_topics: list | None = None) -> list:
    """Full lint over a topic-report build. Returns deduplicated findings."""
    findings: list = []
    scope = [(a.get("_source_topic") or a.get("topic") or "")
             for (a, _r, _p) in (items or [])]

    if deck_text:
        findings += lint_artifact_text(deck_text, kind="deck",
                                       scope_topics=scope,
                                       workspace_topics=workspace_topics)
    if html_text:
        findings += lint_artifact_text(html_text, kind="html",
                                       scope_topics=scope,
                                       workspace_topics=workspace_topics)

    for (assessment, forecast_run, _prior) in (items or []):
        raw = forecast_run.get("raw_output") if isinstance(forecast_run, dict) else None
        if isinstance(raw, str):
            import json
            try:
                raw = json.loads(raw)
                if isinstance(raw, str):
                    raw = json.loads(raw)
            except Exception:
                raw = None
        if isinstance(raw, dict):
            findings += lint_run_content(
                raw, assessment.get("_articles_corpus") or [],
                topic=assessment.get("topic") or "",
            )

    # Deduplicate on (check, detail).
    seen: set = set()
    out: list = []
    for f in findings:
        key = (f["check"], f["detail"])
        if key in seen:
            continue
        seen.add(key)
        out.append(f)
    if out:
        counts: dict = {}
        for f in out:
            counts[f["check"]] = counts.get(f["check"], 0) + 1
        _log.warning("report lint: %d finding(s): %s", len(out), counts)
    return out


def deck_text_from_blob(blob: bytes) -> str:
    """All text frames of a rendered PPTX, for the string checks."""
    from io import BytesIO
    try:
        from pptx import Presentation
        prs = Presentation(BytesIO(blob))
        return "\n".join(sh.text_frame.text for s in prs.slides
                         for sh in s.shapes if sh.has_text_frame)
    except Exception as e:
        _log.warning("report lint: could not read deck blob: %s", e)
        return ""
